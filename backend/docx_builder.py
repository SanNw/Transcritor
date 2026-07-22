"""Geração do arquivo .docx a partir das páginas transcritas."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from ocr import PageResult


def _new_document() -> Document:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    return document


def build_docx(pages: list[PageResult], title: str, output_path: Path, include_page_headings: bool = True) -> None:
    document = _new_document()
    document.add_heading(title, level=0)

    for page in pages:
        if include_page_headings:
            heading = page.heading or f"Página {page.number}"
            if page.used_ocr:
                heading += " (OCR)"
            document.add_heading(heading, level=2)

        text = page.text.strip()
        if not text:
            document.add_paragraph("[Nenhum texto reconhecido nesta página]")
            continue

        for paragraph in text.split("\n\n"):
            paragraph = paragraph.strip()
            if paragraph:
                document.add_paragraph(paragraph)

    document.save(output_path)


def build_docx_from_text(text: str, title: str, output_path: Path) -> None:
    """Gera um .docx de fluxo único a partir de um texto já pronto — usado
    para o resultado de pós-processamento por IA (correção, resumo,
    tradução ou qualquer outra instrução livre), que não é mais organizado
    por página/capítulo do documento original.
    """
    document = _new_document()
    document.add_heading(title, level=0)

    for paragraph in text.strip().split("\n\n"):
        paragraph = paragraph.strip()
        if paragraph:
            document.add_paragraph(paragraph)

    document.save(output_path)
